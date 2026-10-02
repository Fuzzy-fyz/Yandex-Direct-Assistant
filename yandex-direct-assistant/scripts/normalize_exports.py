#!/usr/bin/env python3
"""Read and unify Yandex Direct, Yandex Metrica and CRM exports (CSV / XLSX).

Local and read-only by design:
  * only reads the input files and never modifies them;
  * writes results only into the --out directory;
  * opens no network connections (sockets and DNS are disabled at start-up);
  * contains no secrets and needs no tokens;
  * drops columns that look like personal data (phones, e-mails, names, comments).

Supported report types (auto-detected, override with --type):
  campaigns       Direct report by campaigns (Мастер отчётов / статистика)
  search_queries  Direct search query report (поисковые запросы)
  placements      Direct report by placements / РСЯ площадки
  metrica_goals   Metrica report with goals (достижения целей, целевые визиты)
  crm             CRM export with leads / deals and their statuses

Usage:
  python3 normalize_exports.py FILE [FILE ...] --out DIR [--type TYPE]
                               [--target-cpa RUB] [--sheet NAME] [--top N]

Outputs in DIR:
  summary.md                 human/LLM-readable summary with data-quality warnings
  summary.json               the same in machine-readable form
  <file>.normalized.csv      UTF-8 CSV with canonical column names

Standard library only; uses openpyxl for XLSX when installed, otherwise a
built-in minimal XLSX reader.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import math
import re
import socket
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path


# --------------------------------------------------------------------------
# Safety: no network access from this process.
# --------------------------------------------------------------------------

def _network_blocked(*_args, **_kwargs):
    raise RuntimeError("normalize_exports.py: network access is disabled")


class _NoNetworkSocket(socket.socket):
    def connect(self, *args, **kwargs):  # noqa: D401
        _network_blocked()

    def connect_ex(self, *args, **kwargs):
        _network_blocked()

    def sendto(self, *args, **kwargs):
        _network_blocked()


def disable_network() -> None:
    socket.socket = _NoNetworkSocket  # type: ignore[misc]
    socket.create_connection = _network_blocked  # type: ignore[assignment]
    socket.getaddrinfo = _network_blocked  # type: ignore[assignment]


# --------------------------------------------------------------------------
# Column dictionary
# --------------------------------------------------------------------------

REPORT_TYPES = ["campaigns", "search_queries", "placements", "metrica_goals", "crm"]

TYPE_TITLES = {
    "campaigns": "Директ: отчёт по кампаниям",
    "search_queries": "Директ: поисковые запросы",
    "placements": "Директ: площадки",
    "metrica_goals": "Метрика: цели",
    "crm": "CRM: лиды / сделки",
    "unknown": "Тип не определён",
}

ALIASES: dict[str, list[str]] = {
    "date": ["дата", "день", "date", "дата визита", "дата создания", "дата заявки",
             "дата сделки", "дата обращения", "created", "created at", "created_at", "day"],
    "week": ["неделя", "week"],
    "month": ["месяц", "month"],
    "campaign_id": ["№ кампании", "n кампании", "номер кампании", "id кампании",
                    "campaignid", "campaign id", "campaign_id"],
    "campaign": ["кампания", "название кампании", "campaign", "campaignname",
                 "campaign name", "кампания яндекс.директа", "кампания директа",
                 "рекламная кампания"],
    "ad_group_id": ["№ группы", "id группы", "номер группы", "adgroupid", "adgroup id"],
    "ad_group": ["группа", "группа объявлений", "название группы", "adgroupname",
                 "ad group", "adgroup"],
    "ad_id": ["№ объявления", "id объявления", "номер объявления", "adid", "ad id"],
    "keyword": ["условие показа", "ключевая фраза", "фраза", "ключевое слово",
                "критерий", "criterion", "keyword", "условие показа объявления"],
    "search_query": ["поисковый запрос", "запрос", "query", "search query",
                     "searchquery", "поисковая фраза", "поисковая фраза директа"],
    "placement": ["площадка", "placement", "название площадки", "домен площадки",
                  "сайт площадки"],
    "network": ["тип площадки", "место показа", "тип места показа", "adnetworktype",
                "сеть", "тип показа"],
    "impressions": ["показы", "impressions", "кол-во показов", "количество показов"],
    "clicks": ["клики", "clicks", "кол-во кликов", "количество кликов"],
    "ctr": ["ctr"],
    "cost": ["расход", "расходы", "cost", "затраты", "расход всего", "стоимость кликов"],
    "avg_cpc": ["ср. цена клика", "средняя цена клика", "avgcpc", "avg cpc", "cpc",
                "ср.цена клика"],
    "conversions": ["конверсии", "conversions", "количество конверсий", "кол-во конверсий"],
    "conversion_rate": ["конверсия", "conversion rate", "conversionrate", "cr",
                        "конверсия по цели"],
    "cpa": ["стоимость конверсии", "cpa", "цена конверсии", "costperconversion",
            "цена цели", "стоимость цели"],
    "revenue": ["доход", "выручка", "revenue", "доход по цели", "сумма продажи",
                "сумма сделки", "сумма", "бюджет сделки", "сумма заказа"],
    "profit": ["прибыль", "маржинальная прибыль", "маржа", "profit"],
    "drr": ["дрр", "drr"],
    "roi": ["рентабельность", "roi", "romi"],
    "goal": ["цель", "goal", "название цели", "цель метрики"],
    "goal_id": ["id цели", "goalid", "goal id", "№ цели", "номер цели"],
    "goal_reaches": ["достижения цели", "достижения целей", "goal reaches", "reaches",
                     "достижения"],
    "goal_visits": ["целевые визиты", "goal visits", "визиты с достижением цели"],
    "visits": ["визиты", "visits", "сеансы", "sessions"],
    "users": ["посетители", "users", "пользователи"],
    "bounce_rate": ["отказы", "bounce rate", "bouncerate"],
    "source": ["источник трафика", "источник", "traffic source", "канал"],
    "utm_source": ["utm_source", "utm source"],
    "utm_medium": ["utm_medium", "utm medium"],
    "utm_campaign": ["utm_campaign", "utm campaign"],
    "utm_content": ["utm_content", "utm content"],
    "utm_term": ["utm_term", "utm term"],
    "device": ["тип устройства", "устройство", "device", "devicetype"],
    "region": ["регион", "регион таргетинга", "город", "region", "местоположение"],
    "lead_id": ["id лида", "id заявки", "№ заявки", "id сделки", "№ сделки", "lead id",
                "deal id", "номер заявки", "номер сделки", "id"],
    "lead_status": ["статус", "статус лида", "статус заявки", "статус сделки", "этап",
                    "этап сделки", "stage", "status"],
    "lead_quality": ["качество лида", "качество", "квалификация", "lead quality",
                     "целевой лид"],
    "client_id": ["clientid", "client id", "yandex clientid", "yclid",
                  "идентификатор клиента"],
}

NUMERIC_COLUMNS = {
    "impressions", "clicks", "ctr", "cost", "avg_cpc", "conversions", "conversion_rate",
    "cpa", "revenue", "profit", "drr", "roi", "goal_reaches", "goal_visits", "visits",
    "users", "bounce_rate",
}
ADDITIVE_COLUMNS = ["impressions", "clicks", "cost", "conversions", "revenue", "profit",
                    "goal_reaches", "goal_visits", "visits"]
CANONICAL_ORDER = list(ALIASES.keys())

PII_EXACT = {
    "имя", "фио", "ф.и.о.", "name", "full name", "контакт", "контактное лицо", "клиент",
    "телефон", "phone", "email", "e-mail", "почта", "адрес", "address", "комментарий",
    "комментарии", "comment", "comments", "паспорт", "ip", "ip-адрес", "ip адрес",
    "фамилия", "отчество", "имя клиента", "мобильный", "contact",
}
PII_CONTAINS = ["телефон", "phone", "e-mail", "email", "почт", "фио", "паспорт", "фамили",
                "отчеств", "whatsapp", "telegram", "комментар", "адрес доставки",
                "адрес клиента", "имя клиента", "contact name"]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-zа-я]{2,}$", re.I)
PHONE_RE = re.compile(r"^(\+7|8|7)[\s\-(]*\d{3}[\s\-)]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}$")

TOTAL_RE = re.compile(r"^\s*(итого|всего|total|итог|итого и средние|итого/средние)\s*[:.]?\s*$", re.I)
EMPTY_TOKENS = {"", "-", "--", "—", "–", "n/a", "na", "нет данных", "null", "none"}
PERIOD_RE = re.compile(r"(\d{2}\.\d{2}\.\d{4})\s*[-–—]\s*(\d{2}\.\d{2}\.\d{4})")
SUSPICIOUS_ROW_COUNTS = {1000, 5000, 10000, 50000, 65535, 100000, 1048575}

_INVALID = object()


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def norm_text(value) -> str:
    s = "" if value is None else str(value)
    s = s.replace(" ", " ").replace(" ", " ").replace("ё", "е").replace("Ё", "Е")
    return re.sub(r"\s+", " ", s).strip()


def header_key(value) -> str:
    s = norm_text(value).lower()
    s = re.sub(r"\(.*?\)", "", s)
    s = re.sub(r",\s*(руб|₽|rub|%|тенге|usd|eur|у\.е).*$", "", s)
    s = s.replace("₽", "")
    s = re.sub(r"\s+", " ", s).strip(" .,:;")
    return s


ALIAS_LOOKUP: dict[str, str] = {}
for _canon, _names in ALIASES.items():
    for _name in _names:
        ALIAS_LOOKUP.setdefault(header_key(_name), _canon)


def canonical_for(header) -> str | None:
    key = header_key(header)
    if not key:
        return None
    if key in ALIAS_LOOKUP:
        return ALIAS_LOOKUP[key]
    if key.startswith("расход"):
        return "cost"
    if key.startswith("конверсии"):
        return "conversions"
    if key.startswith("доход"):
        return "revenue"
    if key.startswith("достижения цел"):
        return "goal_reaches"
    if key.startswith("целевые визиты"):
        return "goal_visits"
    if key.startswith("конверсия по цели") or key.startswith("конверсия"):
        return "conversion_rate"
    return None


def parenthetical(header) -> str:
    m = re.search(r"\((.*?)\)", norm_text(header)) or re.search(r"«(.*?)»", norm_text(header))
    return m.group(1).strip() if m else ""


def is_pii_header(header) -> bool:
    key = header_key(header)
    if key in PII_EXACT:
        return True
    return any(part in key for part in PII_CONTAINS)


def parse_number(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(value)
    s = norm_text(value)
    if s.lower() in EMPTY_TOKENS:
        return None
    s = re.sub(r"(руб\.?|₽|rub|%|р\.)", "", s, flags=re.I)
    s = s.replace(" ", "").replace("'", "")
    if s.count(",") == 1 and "." not in s:
        s = s.replace(",", ".")
    elif "," in s and s.count(".") == 1:
        s = s.replace(",", "")
    elif s.count(",") > 1 and "." not in s:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return _INVALID


DATE_FORMATS = ["%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%Y.%m.%d", "%d/%m/%Y",
                "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M",
                "%Y-%m-%dT%H:%M:%S"]


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and 20000 < float(value) < 80000:
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(float(value)))
    s = norm_text(value)
    if not s or s.lower() in EMPTY_TOKENS:
        return None
    if re.fullmatch(r"\d{5}(\.0+)?", s):
        return parse_date(float(s))
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(s[:19], fmt).date()
        except ValueError:
            continue
    return _INVALID


def fmt_num(x, digits=2) -> str:
    if x is None:
        return "нет данных"
    if isinstance(x, float) and (math.isnan(x) or math.isinf(x)):
        return "нет данных"
    if abs(x - round(x)) < 1e-9 and digits:
        digits = 0
    s = f"{x:,.{digits}f}".replace(",", " ").replace(".", ",")
    return s


def ratio(a, b, mult=1.0):
    if a is None or b in (None, 0):
        return None
    return a / b * mult


def md_escape(s: str) -> str:
    return norm_text(s).replace("|", "\\|")


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------

def detect_encoding(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return "utf-16"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        return "cp1251"


def detect_delimiter(lines: list[str]) -> str:
    best, best_score = ";", -1
    for delim in [";", "\t", ","]:
        counts = [line.count(delim) for line in lines if line.strip()]
        counts = [c for c in counts if c > 0]
        if not counts:
            continue
        common, freq = Counter(counts).most_common(1)[0]
        score = freq * 1000 + common
        if score > best_score:
            best, best_score = delim, score
    return best


def read_csv(path: Path) -> tuple[list[list], dict]:
    raw = path.read_bytes()
    enc = detect_encoding(raw)
    text = raw.decode(enc, errors="replace")
    lines = text.splitlines()
    meta = {"encoding": enc}
    if lines and lines[0].lower().startswith("sep="):
        delim = lines[0][4:5] or ";"
        text = "\n".join(lines[1:])
    else:
        delim = detect_delimiter(lines[:40])
    meta["delimiter"] = {"\t": "tab"}.get(delim, delim)
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
    return rows, meta


XLSX_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _col_index(ref: str) -> int:
    letters = re.match(r"[A-Z]+", ref or "A").group(0)
    idx = 0
    for ch in letters:
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


def read_xlsx_builtin(path: Path, sheet: str | None) -> tuple[list[list], dict]:
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root.iter(XLSX_NS + "si"):
                shared.append("".join(t.text or "" for t in si.iter(XLSX_NS + "t")))
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        sheets = [(s.get("name"), s.get(rel_ns)) for s in wb.iter(XLSX_NS + "sheet")]
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        relmap = {r.get("Id"): r.get("Target") for r in rels}
        chosen = sheets[0]
        if sheet:
            matches = [s for s in sheets if s[0] == sheet]
            if not matches:
                raise ValueError(f"лист «{sheet}» не найден; есть: {[s[0] for s in sheets]}")
            chosen = matches[0]
        target = relmap[chosen[1]]
        target = target.lstrip("/") if target.startswith("/") else "xl/" + target
        root = ET.fromstring(z.read(target))
        rows: list[list] = []
        for row in root.iter(XLSX_NS + "row"):
            cells: dict[int, object] = {}
            for c in row.findall(XLSX_NS + "c"):
                col = _col_index(c.get("r", "A"))
                t = c.get("t")
                v = c.find(XLSX_NS + "v")
                if t == "s" and v is not None:
                    val = shared[int(v.text)]
                elif t == "inlineStr":
                    val = "".join(x.text or "" for x in c.iter(XLSX_NS + "t"))
                elif t in ("str", "e") and v is not None:
                    val = v.text
                elif t == "b" and v is not None:
                    val = v.text == "1"
                elif v is not None and v.text not in (None, ""):
                    try:
                        val = float(v.text)
                    except ValueError:
                        val = v.text
                else:
                    val = None
                cells[col] = val
            rows.append([cells.get(i) for i in range(max(cells) + 1)] if cells else [])
    return rows, {"sheet": chosen[0], "sheets": [s[0] for s in sheets], "reader": "builtin"}


def read_xlsx(path: Path, sheet: str | None) -> tuple[list[list], dict]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return read_xlsx_builtin(path, sheet)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        if sheet and sheet not in names:
            raise ValueError(f"лист «{sheet}» не найден; есть: {names}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        return rows, {"sheet": ws.title, "sheets": names, "reader": "openpyxl"}
    finally:
        wb.close()


def read_table(path: Path, sheet: str | None) -> tuple[list[list], dict]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path, sheet)
    if ext in (".csv", ".tsv", ".txt"):
        return read_csv(path)
    if ext == ".xls":
        raise ValueError("старый формат .xls не поддерживается — сохраните файл как XLSX или CSV")
    raise ValueError(f"неподдерживаемый формат {ext}; нужен CSV или XLSX")


# --------------------------------------------------------------------------
# Normalisation
# --------------------------------------------------------------------------

def find_header(rows: list[list]) -> int | None:
    best_idx, best_score = None, 0
    for idx, row in enumerate(rows[:25]):
        mapped = {canonical_for(c) for c in row if norm_text(c)}
        mapped.discard(None)
        if len(mapped) > best_score:
            best_idx, best_score = idx, len(mapped)
    return best_idx if best_score >= 2 else None


def detect_type(columns: set[str]) -> str:
    traffic = columns & {"impressions", "clicks", "cost"}
    if "search_query" in columns and traffic:
        return "search_queries"
    if "placement" in columns and traffic:
        return "placements"
    if not columns & {"impressions", "clicks"} and columns & {"lead_status", "lead_quality", "lead_id"}:
        return "crm"
    if columns & {"goal_reaches", "goal_visits", "visits", "goal"} and "impressions" not in columns:
        return "metrica_goals"
    if columns & {"campaign", "campaign_id"} and traffic:
        return "campaigns"
    return "unknown"


def first_text(row: list) -> str:
    for cell in row[:4]:
        s = norm_text(cell)
        if s:
            return s
    return ""


def normalise_file(path: Path, args) -> dict:
    result: dict = {"file": path.name, "warnings": [], "info": [], "errors": []}
    warn, info, err = result["warnings"].append, result["info"].append, result["errors"].append
    try:
        rows, meta = read_table(path, args.sheet)
    except Exception as exc:  # noqa: BLE001 - report any parsing problem to the user
        err(f"не удалось прочитать файл: {exc}")
        return result
    result["source"] = meta
    if meta.get("sheets") and len(meta["sheets"]) > 1:
        info(f"в книге несколько листов {meta['sheets']}; обработан «{meta['sheet']}» "
             f"(другой лист — параметр --sheet)")

    h = find_header(rows)
    if h is None:
        err("не найдена строка заголовков со знакомыми столбцами (Кампания, Расход, Клики, "
            "Поисковый запрос, Площадка, Цель …)")
        return result

    meta_lines = [first_text(r) for r in rows[:h] if first_text(r)]
    result["meta_lines"] = meta_lines
    header = [norm_text(c) for c in rows[h]]

    # map columns
    mapping: list[tuple[int, str, str]] = []  # (index, output name, original header)
    used: Counter = Counter()
    dropped_pii: list[str] = []
    duplicate_groups: dict[str, list[str]] = defaultdict(list)
    for idx, name in enumerate(header):
        if not name:
            continue
        if is_pii_header(name):
            dropped_pii.append(name)
            continue
        canon = canonical_for(name)
        if canon:
            duplicate_groups[canon].append(name)
            if used[canon]:
                label = parenthetical(name) or str(used[canon] + 1)
                out_name = f"{canon}[{label}]"
            else:
                out_name = canon
            used[canon] += 1
            mapping.append((idx, out_name, name))
        else:
            mapping.append((idx, name, name))

    canon_present = {m[1] for m in mapping if m[1] in ALIASES}
    rtype = args.type or detect_type(canon_present)
    result["report_type"] = rtype
    result["report_title"] = TYPE_TITLES.get(rtype, rtype)
    result["recognised_columns"] = OrderedDict((m[2], m[1]) for m in mapping if m[1] != m[2])
    result["unrecognised_columns"] = [m[2] for m in mapping if m[1] == m[2]]
    for canon, names in duplicate_groups.items():
        if len(names) > 1:
            warn(f"несколько столбцов типа «{canon}»: {names}. Основным взят первый "
                 f"(«{names[0]}»); не суммируйте их — это могут быть разные цели")

    # data rows
    records: list[dict] = []
    totals_rows: list[dict] = []
    bad_numbers: Counter = Counter()
    bad_dates = 0
    for raw in rows[h + 1:]:
        if not any(norm_text(c) for c in raw):
            continue
        label = first_text(raw)
        rec: dict = {}
        for idx, out_name, _orig in mapping:
            rec[out_name] = raw[idx] if idx < len(raw) else None
        if TOTAL_RE.match(label):
            totals_rows.append(rec)
            continue
        non_empty = [c for c in raw if norm_text(c)]
        if len(non_empty) == 1 and len(mapping) > 2 and parse_number(non_empty[0]) in (None, _INVALID):
            meta_lines.append(norm_text(non_empty[0]))  # footer line such as "Отчёт сформирован …"
            continue
        records.append(rec)

    # value-level PII scan on unrecognised text columns
    for _idx, out_name, orig in list(mapping):
        if out_name != orig:
            continue
        values = [norm_text(r.get(out_name)) for r in records if norm_text(r.get(out_name))]
        if values:
            hits = sum(1 for v in values if EMAIL_RE.match(v) or PHONE_RE.match(v))
            if hits / len(values) > 0.3:
                dropped_pii.append(orig)
                mapping.remove((_idx, out_name, orig))
                for r in records:
                    r.pop(out_name, None)
    if dropped_pii:
        info(f"удалены столбцы, похожие на персональные данные: {dropped_pii}")
    result["dropped_pii_columns"] = dropped_pii

    # typed values
    for rec in records:
        for key in list(rec.keys()):
            base = key.split("[")[0]
            if base in NUMERIC_COLUMNS:
                val = parse_number(rec[key])
                if val is _INVALID:
                    bad_numbers[key] += 1
                    val = None
                rec[key] = val
            elif base == "date":
                d = parse_date(rec[key])
                if d is _INVALID:
                    bad_dates += 1
                    d = None
                rec[key] = d.isoformat() if d else None
            else:
                rec[key] = norm_text(rec[key])
    if bad_numbers:
        warn(f"не распознаны числовые значения: {dict(bad_numbers)} — они пропущены")
    if bad_dates:
        warn(f"не распознано дат: {bad_dates}")

    result["rows"] = len(records)
    result["totals_rows_removed"] = len(totals_rows)
    result["_records"] = records
    result["_columns"] = [m[1] for m in mapping]

    if not records:
        err("нет строк с данными после заголовка")
        return result

    run_quality_checks(result, records, totals_rows, canon_present, rtype, meta_lines, header)
    build_aggregates(result, records, rtype, args)
    return result


# --------------------------------------------------------------------------
# Quality checks
# --------------------------------------------------------------------------

def col_sum(records, key):
    vals = [r.get(key) for r in records if isinstance(r.get(key), (int, float))]
    return sum(vals) if vals else None


def run_quality_checks(result, records, totals_rows, cols, rtype, meta_lines, header):
    warn, info = result["warnings"].append, result["info"].append
    meta_text = " ".join(meta_lines + header).lower()

    # period
    dates = sorted({r["date"] for r in records if r.get("date")})
    title_period = PERIOD_RE.search(" ".join(meta_lines))
    if dates:
        result["period"] = {"from": dates[0], "to": dates[-1], "source": "столбец дат"}
        d0, d1 = dt.date.fromisoformat(dates[0]), dt.date.fromisoformat(dates[-1])
        span = (d1 - d0).days + 1
        if span <= 400 and len(dates) > 1 and rtype != "crm":
            present = set(dates)
            missing = [(d0 + dt.timedelta(days=i)).isoformat() for i in range(span)
                       if (d0 + dt.timedelta(days=i)).isoformat() not in present]
            if missing:
                warn(f"разрыв во временном ряду: нет данных за {len(missing)} дн. "
                     f"(например, {', '.join(missing[:5])}) — проверьте остановки показов "
                     f"или неполную выгрузку")
        if d1 > dt.date.today():
            warn(f"в данных есть даты из будущего ({dates[-1]}) — проверьте формат дат")
    elif title_period:
        tp = [dt.datetime.strptime(x, "%d.%m.%Y").date().isoformat() for x in title_period.groups()]
        result["period"] = {"from": tp[0], "to": tp[1], "source": "заголовок отчёта"}
    else:
        result["period"] = None
        warn("период не найден ни в столбце дат, ни в заголовке — уточните период выгрузки")
    if dates and title_period:
        info(f"период в заголовке отчёта: {title_period.group(1)} — {title_period.group(2)}")

    # duplicates
    keyed = Counter(tuple(sorted((k, str(v)) for k, v in r.items())) for r in records)
    dups = sum(c - 1 for c in keyed.values() if c > 1)
    if dups:
        warn(f"полностью повторяющихся строк: {dups} — возможна двойная выгрузка")

    # negatives
    neg = Counter(k for r in records for k, v in r.items()
                  if isinstance(v, float) and v < 0 and k.split("[")[0] in NUMERIC_COLUMNS
                  and k.split("[")[0] not in ("roi", "profit"))
    if neg:
        warn(f"отрицательные значения: {dict(neg)}")

    # totals row cross-check
    for tot in totals_rows[:1]:
        for key in ADDITIVE_COLUMNS:
            if key not in tot:
                continue
            t = parse_number(tot.get(key))
            s = col_sum(records, key)
            if isinstance(t, float) and s is not None and t and abs(s - t) / abs(t) > 0.01:
                warn(f"сумма строк по «{key}» ({fmt_num(s)}) не совпадает с итоговой строкой "
                     f"({fmt_num(t)}) — выгрузка может быть обрезана или отфильтрована")
    if totals_rows:
        info(f"удалено итоговых строк: {len(totals_rows)} (в расчётах не участвуют)")

    if len(records) in SUSPICIOUS_ROW_COUNTS:
        warn(f"ровно {len(records)} строк — похоже на лимит выгрузки, данные могут быть неполными")

    direct = rtype in ("campaigns", "search_queries", "placements")
    if direct:
        if "cost" not in cols:
            warn("нет столбца расхода — экономику (CPA, ДРР) посчитать нельзя")
        if "conversions" not in cols:
            warn("нет столбца конверсий — в Мастере отчётов выберите цели Метрики; без них "
                 "оценка возможна только по кликам, что недостаточно для решений")
        else:
            conv = col_sum(records, "conversions") or 0
            cost = col_sum(records, "cost") or 0
            if conv == 0 and cost > 0:
                warn("конверсий 0 при ненулевом расходе — проверьте, выбраны ли цели в отчёте "
                     "и связаны ли Директ и Метрика")
            over = sum(1 for r in records if (r.get("conversions") or 0) > (r.get("clicks") or 0) > 0)
            if over:
                info(f"строк, где конверсий больше, чем кликов: {over} — вероятно, учитывается "
                     f"несколько целей или несколько достижений за визит")
        bad_ctr = sum(1 for r in records if (r.get("clicks") or 0) > (r.get("impressions") or math.inf))
        if bad_ctr:
            warn(f"строк, где кликов больше, чем показов: {bad_ctr} — проверьте выгрузку")
        if "revenue" not in cols:
            info("нет дохода/выручки — ДРР, ROAS и ROMI недоступны; нужны данные CRM "
                 "или ценность целей")
        if "ндс" in meta_text:
            vat = "без НДС" if "без ндс" in meta_text else "с НДС"
            info(f"в выгрузке указано: расход {vat}")
        else:
            info("не указано, с НДС ли расход — уточните перед сравнением с бюджетом")
        if "атрибуц" in meta_text:
            m = re.search(r"[^.;]*атрибуц[^.;]*", meta_text)
            info(f"модель атрибуции из заголовка: «{m.group(0).strip() if m else 'указана'}»")
        else:
            info("модель атрибуции в выгрузке не указана — уточните её при сравнении с Метрикой")
        if rtype == "search_queries" and "keyword" not in cols:
            info("в отчёте нет условия показа (ключевой фразы) — связать запросы с фразами нельзя")

    if rtype == "metrica_goals":
        goals = {r.get("goal") for r in records if r.get("goal")}
        goal_cols = [c for c in result["_columns"] if c.split("[")[0] in ("goal_reaches", "goal_visits", "conversions")]
        if len(goals) > 1:
            warn(f"целей в файле: {len(goals)} — не суммируйте достижения и визиты разных целей "
                 f"(двойной учёт одного обращения)")
        if len(goal_cols) > 1:
            info(f"несколько метрик целей в файле: {goal_cols}; «достижения» и «целевые визиты» "
                 f"считаются по-разному — для CPA выберите одну")
        if not goals and not goal_cols:
            warn("не найдено ни столбца «Цель», ни метрик достижений — проверьте, что в отчёт "
                 "Метрики добавлены цели")
        if "campaign" not in cols and "utm_campaign" not in cols and "source" not in cols:
            info("нет разбивки по кампаниям/источникам — сопоставить с Директом можно только в целом")

    if rtype == "crm":
        if "lead_status" not in cols:
            warn("нет статусов лидов — качество лидов и продажи оценить нельзя")
        if not cols & {"campaign", "utm_campaign", "utm_source", "source", "client_id"}:
            warn("нет UTM-меток, кампании или ClientID — лиды нельзя связать с рекламой")
        if "lead_id" in cols:
            ids = Counter(r.get("lead_id") for r in records if r.get("lead_id"))
            d = sum(c - 1 for c in ids.values() if c > 1)
            if d:
                warn(f"повторяющихся ID лидов/сделок: {d}")

    if rtype == "unknown":
        warn("тип отчёта не определён — укажите его параметром --type; расчёты ниже общие")


# --------------------------------------------------------------------------
# Aggregation
# --------------------------------------------------------------------------

ENTITY_KEYS = {
    "campaigns": ["campaign", "campaign_id"],
    "search_queries": ["search_query"],
    "placements": ["placement"],
    "metrica_goals": ["goal", "campaign", "utm_campaign", "source"],
    "crm": ["utm_campaign", "campaign", "lead_status"],
    "unknown": ["campaign", "campaign_id"],
}


def derived(m: dict) -> dict:
    return {
        "ctr_pct": ratio(m.get("clicks"), m.get("impressions"), 100),
        "cpc": ratio(m.get("cost"), m.get("clicks")),
        "cr_pct": ratio(m.get("conversions"), m.get("clicks"), 100),
        "cpa": ratio(m.get("cost"), m.get("conversions")),
        "drr_pct": ratio(m.get("cost"), m.get("revenue"), 100),
        "roas": ratio(m.get("revenue"), m.get("cost")),
    }


def low_data_flag(cost, conv, ref_cpa):
    if ref_cpa in (None, 0) or cost is None:
        return None, "нет ориентира CPA"
    expected = cost / ref_cpa
    if conv:
        if conv < 3:
            return expected, "конверсий < 3 — CPA неустойчив"
        return expected, ""
    if expected < 1:
        return expected, "мало данных: 0 конв. ещё ничего не доказывает"
    if expected < 2:
        return expected, "пограничный объём"
    return expected, "сигнал: 0 конв. при ожидаемых ≥ 2"


def build_aggregates(result, records, rtype, args):
    cols = set(result["_columns"])
    totals = {k: col_sum(records, k) for k in ADDITIVE_COLUMNS if k in cols}
    if rtype == "metrica_goals":
        n_goals = len({r.get("goal") for r in records if r.get("goal")})
        if n_goals > 1:
            # summing reaches of different goals double-counts the same visitor
            for k in ("goal_reaches", "goal_visits", "conversions", "visits"):
                totals.pop(k, None)
        elif "conversions" not in totals:
            totals["conversions"] = totals.get("goal_reaches", totals.get("goal_visits"))
    totals.update(derived(totals))
    result["totals"] = totals

    ref_cpa, ref_src = None, None
    if args.target_cpa:
        ref_cpa, ref_src = args.target_cpa, "целевой CPA (--target-cpa)"
    elif totals.get("cpa"):
        ref_cpa, ref_src = totals["cpa"], "средний CPA файла"
    result["reference_cpa"] = {"value": ref_cpa, "source": ref_src}

    if rtype == "crm":
        statuses = Counter(r.get("lead_status") or "без статуса" for r in records)
        result["lead_statuses"] = statuses.most_common()

    if rtype == "metrica_goals" and "goal" in cols:
        second = next((k for k in ("campaign", "utm_campaign", "source") if k in cols), None)
        if second:
            pairs: dict = OrderedDict()
            for r in records:
                key = (r.get("goal") or "(пусто)", r.get(second) or "(пусто)")
                row = pairs.setdefault(key, {"a": key[0], "b": key[1]})
                for k in ("goal_reaches", "goal_visits", "visits"):
                    if isinstance(r.get(k), float):
                        row[k] = row.get(k, 0.0) + r[k]
            result["breakdown"] = {"keys": ["goal", second], "rows": list(pairs.values())}

    key_name = next((k for k in ENTITY_KEYS.get(rtype, []) if k in cols), None)
    if not key_name:
        result["entities"] = []
        return
    result["entity_key"] = key_name
    groups: dict[str, dict] = OrderedDict()
    for r in records:
        name = r.get(key_name) or "(пусто)"
        g = groups.setdefault(name, {"name": name, "rows": 0, "campaigns": set(),
                                     "statuses": Counter()})
        g["rows"] += 1
        if rtype == "crm":
            g["statuses"][r.get("lead_status") or "без статуса"] += 1
        if r.get("campaign"):
            g["campaigns"].add(r["campaign"])
        for k in ADDITIVE_COLUMNS:
            v = r.get(k)
            if isinstance(v, float):
                g[k] = g.get(k, 0.0) + v
    entities = []
    for g in groups.values():
        if rtype == "metrica_goals" and "conversions" not in g:
            g["conversions"] = g.get("goal_reaches", g.get("goal_visits"))
        g.update(derived(g))
        g["campaigns"] = len(g["campaigns"])
        g["statuses"] = dict(g["statuses"]) if rtype == "crm" else None
        if g["statuses"] is None:
            g.pop("statuses")
        if "cost" in g and rtype in ("campaigns", "search_queries", "placements", "unknown"):
            g["expected_conversions"], g["flag"] = low_data_flag(g.get("cost"), g.get("conversions"), ref_cpa)
        entities.append(g)
    sort_key = "cost" if "cost" in cols else ("conversions" if rtype != "crm" else "rows")
    entities.sort(key=lambda e: (e.get(sort_key) or 0, e.get("rows")), reverse=True)
    result["entities"] = entities


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

def write_normalised_csv(result: dict, out_dir: Path) -> Path:
    cols = result["_columns"]
    canon = [c for c in CANONICAL_ORDER if c in cols]
    variants = [c for c in cols if "[" in c]
    others = [c for c in cols if c not in canon and c not in variants]
    fieldnames = ["source_file", "report_type"] + canon + variants + others
    out = out_dir / (Path(result["file"]).stem + ".normalized.csv")
    with out.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for r in result["_records"]:
            row = {"source_file": result["file"], "report_type": result["report_type"]}
            for k, v in r.items():
                row[k] = int(v) if isinstance(v, float) and v.is_integer() else v
            w.writerow(row)
    return out


def entity_table(result: dict, top: int) -> list[str]:
    ents = result.get("entities") or []
    if not ents:
        return []
    rtype = result["report_type"]
    lines = []
    if rtype == "crm":
        names = [s for s, _ in (result.get("lead_statuses") or [])][:6]
        lines += ["| Объект | Лидов | " + " | ".join(md_escape(n) for n in names) + " | Выручка |",
                  "|---|---|" + "---|" * len(names) + "---|"]
        for e in ents[:top]:
            st = e.get("statuses") or {}
            lines.append(f"| {md_escape(e['name'])} | {e['rows']} | "
                         + " | ".join(str(st.get(n, 0)) for n in names)
                         + f" | {fmt_num(e.get('revenue'))} |")
        return lines
    has_cost = any("cost" in e for e in ents)
    if not has_cost:
        lines += ["| Объект | Достижения цели | Целевые визиты | Визиты |", "|---|---|---|---|"]
        for e in ents[:top]:
            lines.append(f"| {md_escape(e['name'])} | {fmt_num(e.get('goal_reaches', e.get('conversions')))} | "
                         f"{fmt_num(e.get('goal_visits'))} | {fmt_num(e.get('visits'))} |")
        return lines
    lines += ["| Объект | Расход | Клики | CTR % | Конв. | CPA | CR % | Ожид. конв. | Флаг |",
              "|---|---|---|---|---|---|---|---|---|"]
    for e in ents[:top]:
        lines.append(
            f"| {md_escape(e['name'])} | {fmt_num(e.get('cost'))} | {fmt_num(e.get('clicks'))} | "
            f"{fmt_num(e.get('ctr_pct'))} | {fmt_num(e.get('conversions'))} | {fmt_num(e.get('cpa'))} | "
            f"{fmt_num(e.get('cr_pct'))} | {fmt_num(e.get('expected_conversions'), 1)} | "
            f"{e.get('flag') or ''} |")
    return lines


def zero_conv_table(result: dict, top: int) -> list[str]:
    ents = [e for e in (result.get("entities") or [])
            if "cost" in e and (e.get("cost") or 0) > 0 and not e.get("conversions")
            and "conversions" in result["_columns"]]
    if not ents:
        return []
    lines = ["| Объект | Расход | Клики | Ожид. конв. | Флаг |", "|---|---|---|---|---|"]
    for e in ents[:top]:
        lines.append(f"| {md_escape(e['name'])} | {fmt_num(e.get('cost'))} | {fmt_num(e.get('clicks'))} | "
                     f"{fmt_num(e.get('expected_conversions'), 1)} | {e.get('flag') or ''} |")
    return lines


TOTAL_LABELS = [
    ("impressions", "Показы"), ("clicks", "Клики"), ("cost", "Расход"),
    ("conversions", "Конверсии"), ("goal_reaches", "Достижения целей"),
    ("goal_visits", "Целевые визиты"), ("visits", "Визиты"), ("revenue", "Доход / выручка"),
    ("profit", "Прибыль"), ("ctr_pct", "CTR, %"), ("cpc", "Ср. цена клика"),
    ("cr_pct", "Конверсия, %"), ("cpa", "CPA"), ("drr_pct", "ДРР, %"), ("roas", "ROAS"),
]


def write_summary(results: list[dict], out_dir: Path, args) -> None:
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    md = ["# Сводка по выгрузкам", "",
          f"Сформировано: {now}. Обработка локальная, данные никуда не передавались.", ""]
    periods = {(r["period"]["from"], r["period"]["to"]) for r in results if r.get("period")}
    if len(periods) > 1:
        md += ["**Внимание:** периоды файлов различаются — "
               + "; ".join(f"{a} — {b}" for a, b in sorted(periods))
               + ". Сравнивайте только сопоставимые периоды.", ""]
    for n, r in enumerate(results, 1):
        md.append(f"## Файл {n}: {r['file']}")
        md.append("")
        if r.get("errors"):
            for e in r["errors"]:
                md.append(f"- [ОШИБКА] {e}")
            md.append("")
            continue
        src = r.get("source", {})
        md.append(f"- Тип отчёта: **{r['report_title']}** (`{r['report_type']}`)")
        if src.get("encoding"):
            md.append(f"- Кодировка / разделитель: {src['encoding']} / {src.get('delimiter')}")
        if src.get("sheet"):
            md.append(f"- Лист: {src['sheet']}")
        p = r.get("period")
        md.append(f"- Период: {p['from']} — {p['to']} ({p['source']})" if p else "- Период: не найден")
        md.append(f"- Строк данных: {r.get('rows', 0)}; итоговых строк удалено: {r.get('totals_rows_removed', 0)}")
        if r.get("meta_lines"):
            md.append(f"- Служебные строки выгрузки: {' | '.join(r['meta_lines'][:4])}")
        rec = r.get("recognised_columns") or {}
        md.append("- Распознанные столбцы: " + ", ".join(f"{k} → {v}" for k, v in rec.items()))
        if r.get("unrecognised_columns"):
            md.append("- Нераспознанные столбцы (сохранены как есть): " + ", ".join(r["unrecognised_columns"]))
        if r.get("normalized_csv"):
            md.append(f"- Нормализованный файл: {r['normalized_csv']}")
        md.append("")
        direct = r["report_type"] in ("campaigns", "search_queries", "placements")
        totals = r.get("totals") or {}
        total_rows = []
        for key, label in TOTAL_LABELS:
            if key not in totals:
                continue
            if totals[key] is None and not (direct and key in ("cpa", "drr_pct", "roas")):
                continue
            total_rows.append(f"| {label} | {fmt_num(totals[key])} |")
        ref = r.get("reference_cpa") or {}
        if ref.get("value"):
            total_rows.append(f"| Ориентир CPA для флагов | {fmt_num(ref['value'])} ({ref['source']}) |")
        md += ["### Итоги", ""]
        if total_rows:
            md += ["| Показатель | Значение |", "|---|---|"] + total_rows + [""]
        else:
            md += ["Общие итоги не суммируются (например, в файле несколько целей) — см. разбивку ниже.", ""]
        if r.get("breakdown"):
            b = r["breakdown"]
            md += [f"### Разбивка {b['keys'][0]} × {b['keys'][1]}", "",
                   f"| {b['keys'][0]} | {b['keys'][1]} | Достижения цели | Целевые визиты | Визиты |",
                   "|---|---|---|---|---|"]
            for row in b["rows"][:60]:
                md.append(f"| {md_escape(row['a'])} | {md_escape(row['b'])} | {fmt_num(row.get('goal_reaches'))} | "
                          f"{fmt_num(row.get('goal_visits'))} | {fmt_num(row.get('visits'))} |")
            md.append("")
        if r.get("lead_statuses"):
            md += ["### Статусы лидов", "", "| Статус | Количество |", "|---|---|"]
            md += [f"| {md_escape(s)} | {c} |" for s, c in r["lead_statuses"]]
            md.append("")
        table = entity_table(r, args.top)
        if table:
            md += [f"### Топ-{args.top} по ключу «{r.get('entity_key')}»", ""] + table + [""]
        zc = zero_conv_table(r, 10)
        if zc:
            md += ["### Расход без конверсий (кандидаты на проверку, не на автоматическое исключение)", ""] + zc + [""]
        md += ["### Качество данных", ""]
        for w in r.get("warnings", []):
            md.append(f"- [!] {w}")
        for i in r.get("info", []):
            md.append(f"- [i] {i}")
        if not r.get("warnings") and not r.get("info"):
            md.append("- замечаний нет")
        md.append("")
    md += ["---",
           "Флаги «Ожид. конв.» = расход ÷ ориентир CPA. Это подсказка для оценки объёма данных, "
           "а не решение: окончательный вывод делает аналитик с учётом бизнес-правил и периода."]
    (out_dir / "summary.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    def clean(obj):
        if isinstance(obj, dict):
            return {k: clean(v) for k, v in obj.items() if not k.startswith("_")}
        if isinstance(obj, list):
            return [clean(v) for v in obj]
        if isinstance(obj, float) and (math.isnan(obj) or math.isinf(obj)):
            return None
        return obj

    payload = {"generated_at": now, "files": [clean(r) for r in results]}
    for f in payload["files"]:
        if "entities" in f:
            f["entities"] = f["entities"][:200]
    (out_dir / "summary.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                                          encoding="utf-8")


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main(argv=None) -> int:
    disable_network()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", type=Path, help="CSV/XLSX exports")
    ap.add_argument("--out", required=True, type=Path, help="output directory for results")
    ap.add_argument("--type", choices=REPORT_TYPES, help="force report type for all files")
    ap.add_argument("--target-cpa", type=float, help="target CPA in account currency, for low-data flags")
    ap.add_argument("--sheet", help="XLSX sheet name (default: first sheet)")
    ap.add_argument("--top", type=int, default=15, help="rows in top tables (default 15)")
    args = ap.parse_args(argv)

    inputs = [p.resolve() for p in args.files]
    for p in inputs:
        if not p.is_file():
            ap.error(f"файл не найден: {p}")
    out_dir = args.out.resolve()
    if any(out_dir == p.parent for p in inputs):
        ap.error("--out не должен совпадать с папкой исходных выгрузок: укажите отдельную папку")
    out_dir.mkdir(parents=True, exist_ok=True)

    results = []
    for p in inputs:
        r = normalise_file(p, args)
        if r.get("_records"):
            target = write_normalised_csv(r, out_dir)
            if target.resolve() in inputs:
                raise SystemExit("отказ: выходной файл совпал с исходным")
            r["normalized_csv"] = target.name
        results.append(r)
    write_summary(results, out_dir, args)

    failed = [r["file"] for r in results if r.get("errors")]
    print(f"Готово: {out_dir / 'summary.md'}")
    for r in results:
        status = "ошибка" if r.get("errors") else f"{r.get('report_type')}, строк {r.get('rows', 0)}, " \
                 f"предупреждений {len(r.get('warnings', []))}"
        print(f"  {r['file']}: {status}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
